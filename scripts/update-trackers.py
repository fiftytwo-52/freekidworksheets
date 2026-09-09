import json
import sys
from pathlib import Path
from openpyxl import load_workbook

manifest = Path(sys.argv[1])
tracker_dir = manifest.parent
md_path = tracker_dir / 'worksheet-upload-tracking.md'
xlsx_path = tracker_dir / 'worksheets-tracker.xlsx'
records = json.loads(manifest.read_text())

# Reconcile imported filenames with existing rows, then put the uploaded batch first.
new_names = {r['originalName'] for r in records}
text = md_path.read_text()
lines = text.splitlines()
header_index = next(i for i, line in enumerate(lines) if line.startswith('| Code |'))
separator_index = header_index + 1
existing_rows = []
for line in lines[separator_index + 1:]:
    if not line.startswith('|'):
        existing_rows.append(line)
        continue
    cells = [cell.strip() for cell in line.strip('|').split('|')]
    if len(cells) >= 10 and cells[8] not in new_names:
        existing_rows.append(line)
new_rows = [
    '| {code} | {title} | {language} | {type} | {colour} | {origin} | {age} | {date} | {filename} | {status} |'.format(
        code=r['code'], title=r['title'], language=r['language'], type=r['type'],
        colour=r['colour'], origin=r['origin'], age=r['ageGroup'], date=r['date'],
        filename=r['originalName'], status=r['status'])
    for r in records
]
md_path.write_text('\n'.join(lines[:separator_index + 1] + new_rows + existing_rows) + '\n')

wb = load_workbook(xlsx_path)
ws = wb['Worksheets']
existing = list(ws.iter_rows(min_row=2, values_only=True))
ws.delete_rows(2, ws.max_row)
for r in records:
    ws.append([r['code'], r['title'], r['language'], r['type'], r['colour'], r['origin'], r['ageGroup'], r['date'], r['originalName'], r['status']])
for row in existing:
    if len(row) >= 10 and row[8] not in new_names:
        ws.append(list(row))

summary = wb['Summary']
values = {row[0].value: row[1].value for row in summary.iter_rows(min_row=2, max_col=2) if row[0].value}
all_rows = list(ws.iter_rows(min_row=2, values_only=True))
values['Total worksheets'] = len(all_rows)
values['English (3xxx)'] = sum(1 for r in all_rows if r[2] == 'English')
values['Nepali (4xxx)'] = sum(1 for r in all_rows if r[2] == 'Nepali')
values['Portuguese (5xxxx)'] = sum(1 for r in all_rows if r[2] == 'Portuguese')
for key in ['Original', 'Altered', 'Black-and-white', 'Colorful']:
    if key in ('Black-and-white', 'Colorful'):
        values[key] = sum(1 for r in all_rows if r[4] == key.lower())
    else:
        values[key] = sum(1 for r in all_rows if r[5] == key)
values['Next English code'] = max([int(r[0]) for r in all_rows if r[2] == 'English' and str(r[0]).isdigit()] or [3052]) + 1
values['Next Nepali code'] = max([int(r[0]) for r in all_rows if r[2] == 'Nepali' and str(r[0]).isdigit()] or [4023]) + 1
values['Next Portuguese code'] = max([int(r[0]) for r in all_rows if r[2] == 'Portuguese' and str(r[0]).isdigit()] or [50000]) + 1
values['Unuploaded / pending publication'] = sum(1 for r in all_rows if str(r[9]).lower() != 'uploaded')
for row in summary.iter_rows(min_row=2, max_col=2):
    if row[0].value in values:
        row[1].value = values[row[0].value]
for key, value in values.items():
    if not any(row[0].value == key for row in summary.iter_rows(min_row=2, max_col=1)):
        summary.append([key, value])
wb.save(xlsx_path)
manifest.unlink()
print(f'Updated trackers with {len(records)} uploaded records.')
